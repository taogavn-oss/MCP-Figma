#!/usr/bin/env python3
"""
doc-generator: Markdown to Word (.docx) Converter
Converts Markdown documents into beautifully styled Microsoft Word documents (.docx).
Supports tables, code blocks, callouts, headings, lists, slide page breaks, and custom themes.
"""

import os
import sys
import re
import argparse
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

THEMES = {
    "modern-blue": {
        "h1": RGBColor(0x0F, 0x17, 0x2A),
        "h2": RGBColor(0x1E, 0x3A, 0x8A),
        "h3": RGBColor(0x25, 0x63, 0xEB),
        "table_header_bg": "1E3A8A",
        "table_header_fg": RGBColor(0xFF, 0xFF, 0xFF),
        "quote_fg": RGBColor(0x47, 0x55, 0x69),
        "badge_bg": "EFF6FF",
    },
    "warm-artisan": {
        "h1": RGBColor(0x1C, 0x19, 0x17),
        "h2": RGBColor(0xDA, 0x29, 0x1C),
        "h3": RGBColor(0xC5, 0x9B, 0x58),
        "table_header_bg": "1C1917",
        "table_header_fg": RGBColor(0xFF, 0xFF, 0xFF),
        "quote_fg": RGBColor(0x78, 0x71, 0x6C),
        "badge_bg": "FEF2F2",
    },
    "slate-minimal": {
        "h1": RGBColor(0x11, 0x18, 0x27),
        "h2": RGBColor(0x37, 0x41, 0x51),
        "h3": RGBColor(0x4B, 0x55, 0x63),
        "table_header_bg": "1F2937",
        "table_header_fg": RGBColor(0xFF, 0xFF, 0xFF),
        "quote_fg": RGBColor(0x6B, 0x72, 0x80),
        "badge_bg": "F3F4F6",
    }
}

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def convert_md_to_docx(md_path, docx_path, theme_name="modern-blue", slide_breaks=True):
    theme = THEMES.get(theme_name, THEMES["modern-blue"])
    doc = Document()

    # Document margins: 0.8 in top/bottom, 0.9 in left/right
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Base Normal font
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Segoe UI'
    font.size = Pt(10.5)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        
        valid_rows = []
        for r in table_rows:
            cells = [c.strip() for c in r.strip().strip('|').split('|')]
            if all(set(c).issubset({'-', ':', ' '}) for c in cells):
                continue
            valid_rows.append(cells)

        if not valid_rows:
            in_table = False
            table_rows = []
            return

        col_count = max(len(r) for r in valid_rows)
        tbl = doc.add_table(rows=len(valid_rows), cols=col_count)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = True

        for i, row in enumerate(valid_rows):
            for j, val in enumerate(row):
                if j < col_count:
                    cell = tbl.cell(i, j)
                    cell.text = val
                    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    if i == 0:
                        set_cell_background(cell, theme["table_header_bg"])
                        for run in p.runs:
                            run.font.bold = True
                            run.font.color.rgb = theme["table_header_fg"]
                            run.font.size = Pt(9.5)
                    else:
                        bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
                        set_cell_background(cell, bg)
                        for run in p.runs:
                            run.font.size = Pt(9.5)
                            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        in_table = False
        table_rows = []

    for line in lines:
        raw = line.rstrip('\r\n')
        stripped = raw.strip()

        # Slide page break
        if stripped.startswith('<!-- slide -->'):
            if in_table:
                flush_table()
            if slide_breaks:
                doc.add_page_break()
            continue

        # Code block
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run('\n'.join(code_lines))
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
                code_lines = []
            else:
                if in_table:
                    flush_table()
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(raw)
            continue

        # Markdown tables
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_rows.append(stripped)
            continue
        elif in_table:
            flush_table()

        # Headings
        if stripped.startswith('# '):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(6)
            run = h.add_run(stripped[2:])
            run.font.name = 'Segoe UI'
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = theme["h1"]
        elif stripped.startswith('## '):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            run = h.add_run(stripped[3:])
            run.font.name = 'Segoe UI'
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = theme["h2"]
        elif stripped.startswith('### '):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            run = h.add_run(stripped[4:])
            run.font.name = 'Segoe UI'
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.color.rgb = theme["h3"]
        elif stripped.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(stripped[5:])
            run.font.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        elif stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped[2:])
            run.font.italic = True
            run.font.color.rgb = theme["quote_fg"]
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            text = stripped[2:]
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)
        elif re.match(r'^\d+\.\s', stripped):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            text = re.sub(r'^\d+\.\s', '', stripped)
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)
        elif stripped == '---':
            pass
        elif stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            parts = re.split(r'(\*\*.*?\*\*)', stripped)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)

    if in_table:
        flush_table()

    # Create destination directory if needed
    os.makedirs(os.path.dirname(os.path.abspath(docx_path)), exist_ok=True)
    doc.save(docx_path)
    return docx_path

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="Convert Markdown to styled DOCX document")
    parser.add_argument("input", help="Path to input markdown file")
    parser.add_argument("-o", "--output", help="Path to output .docx file (optional)")
    parser.add_argument("-t", "--theme", default="modern-blue", choices=["modern-blue", "warm-artisan", "slate-minimal"], help="Color theme for headings and tables")
    parser.add_argument("--no-slide-breaks", action="store_true", help="Do not insert page breaks on <!-- slide -->")

    args = parser.parse_args()
    input_file = args.input
    output_file = args.output or os.path.splitext(input_file)[0] + ".docx"
    slide_breaks = not args.no_slide_breaks

    try:
        res = convert_md_to_docx(input_file, output_file, theme_name=args.theme, slide_breaks=slide_breaks)
        print(f"Successfully generated DOCX: {res}")
    except Exception as e:
        print(f"Error generating DOCX: {e}", file=sys.stderr)
        sys.exit(1)
